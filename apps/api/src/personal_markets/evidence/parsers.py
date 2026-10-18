"""Deterministic parsers for user-supplied research and measurement files."""

from __future__ import annotations

import csv
import io
from typing import Any

from personal_markets.integrations.csv_import import preview_csv

MAX_ROWS = 20_000
MAX_PDF_PAGES = 40
MAX_TEXT_CHARS = 200_000
MAX_SHEETS = 10

ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".pdf", ".txt", ".md"}
_PDF_MAGIC = b"%PDF"
_XLSX_MAGIC = b"PK"


class ParseError(ValueError):
    pass


def sniff_kind(filename: str, content: bytes) -> str:
    name = filename.lower()
    if content.startswith(_PDF_MAGIC) or name.endswith(".pdf"):
        if not content.startswith(_PDF_MAGIC):
            raise ParseError("PDF extension does not match file contents")
        return "pdf"
    if name.endswith(".xlsx") or (content.startswith(_XLSX_MAGIC) and name.endswith(".xlsx")):
        if not content.startswith(_XLSX_MAGIC):
            raise ParseError("Excel extension does not match file contents")
        return "xlsx"
    if name.endswith(".csv") or name.endswith(".txt") or name.endswith(".md"):
        return "text"
    raise ParseError("Unsupported file type. Use CSV, XLSX, PDF, or text.")


def parse_upload(filename: str, content: bytes) -> dict[str, Any]:
    kind = sniff_kind(filename, content)
    if kind == "pdf":
        return {"kind": "pdf", **parse_pdf(content)}
    if kind == "xlsx":
        return {"kind": "xlsx", **parse_xlsx(content)}
    text = content.decode("utf-8-sig", errors="replace")
    first_line = text.splitlines()[0] if text.strip() else ""
    if filename.lower().endswith(".csv") or "," in first_line:
        try:
            preview = preview_csv(text, limit=8)
            return {
                "kind": "csv",
                "columns": preview.columns,
                "sample_rows": preview.sample_rows,
                "row_count": preview.row_count,
                "text_excerpt": text[:4000],
                "ambiguous": False,
            }
        except ValueError as error:
            return {
                "kind": "text",
                "text_excerpt": text[:4000],
                "columns": [],
                "ambiguous": True,
                "error": str(error),
            }
    return {
        "kind": "text",
        "text_excerpt": text[:MAX_TEXT_CHARS],
        "columns": [],
        "ambiguous": False,
    }


def parse_xlsx(content: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    buffer = io.BytesIO(content)
    workbook = load_workbook(buffer, read_only=True, data_only=True)
    sheets = workbook.sheetnames[:MAX_SHEETS]
    if not sheets:
        raise ParseError("Workbook has no sheets")
    worksheet = workbook[sheets[0]]
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as error:
        raise ParseError("Spreadsheet is empty") from error
    columns = [
        str(item).strip() if item is not None else f"column_{index}"
        for index, item in enumerate(header_row, start=1)
    ]
    sample: list[dict[str, str]] = []
    total = 0
    for row in rows:
        total += 1
        if total > MAX_ROWS:
            raise ParseError(f"Spreadsheet exceeds {MAX_ROWS} rows")
        if len(sample) < 8:
            sample.append(
                {
                    columns[index]: "" if value is None else str(value)
                    for index, value in enumerate(row)
                    if index < len(columns)
                }
            )
    return {
        "columns": columns,
        "sample_rows": sample,
        "row_count": total,
        "sheets": sheets,
        "ambiguous": False,
        "text_excerpt": "",
    }


def parse_pdf(content: bytes) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ParseError("PDF support is not installed") from error
    reader = PdfReader(io.BytesIO(content))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ParseError(f"PDF exceeds {MAX_PDF_PAGES} pages")
    pages: list[str] = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        pages.append(page.extract_text() or "")
    text = "\n\n".join(pages)[:MAX_TEXT_CHARS]
    return {
        "columns": [],
        "sample_rows": [],
        "row_count": 0,
        "page_count": len(reader.pages),
        "text_excerpt": text[:8000],
        "ambiguous": True,
        "note": (
            "PDF text is untrusted content. "
            "Map quantitative claims before they become evidence."
        ),
    }


def rows_from_table(content: bytes, filename: str, mapping: dict[str, str]) -> list[dict[str, Any]]:
    kind = sniff_kind(filename, content)
    if kind == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        header = [str(item).strip() if item is not None else "" for item in next(rows)]
        records = []
        for raw in rows:
            records.append(
                {
                    header[i]: "" if raw[i] is None else str(raw[i])
                    for i in range(min(len(header), len(raw)))
                }
            )
        return _map_rows(records, mapping)
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    records = [{key: (value or "").strip() for key, value in row.items() if key} for row in reader]
    return _map_rows(records, mapping)


def _map_rows(records: list[dict[str, str]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    metric_col = mapping.get("metric_column")
    value_col = mapping.get("value_column")
    time_col = mapping.get("observed_at_column")
    unit_col = mapping.get("unit_column")
    metric_name = mapping.get("metric") or "uploaded_metric"
    if value_col is None:
        raise ParseError("A value column mapping is required")
    events: list[dict[str, Any]] = []
    for index, row in enumerate(records, start=2):
        raw_value = (row.get(value_col) or "").strip()
        if not raw_value:
            continue
        try:
            numeric = float(raw_value.replace(",", ""))
        except ValueError as error:
            raise ParseError(f"row {index} has a non-numeric value") from error
        events.append(
            {
                "event_type": mapping.get("event_type") or "uploaded_measurement",
                "domain": mapping.get("domain") or "user_state",
                "metric": (row.get(metric_col) or metric_name).strip()
                if metric_col
                else metric_name,
                "numeric_value": numeric,
                "unit": (row.get(unit_col) if unit_col else None) or mapping.get("unit"),
                "observed_at": (row.get(time_col) or "").strip() or "unknown",
                "source_type": "upload",
                "summary": f"Uploaded {metric_name}",
                "observation_kind": mapping.get("observation_kind") or "snapshot",
            }
        )
    return events
